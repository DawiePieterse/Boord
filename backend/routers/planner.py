"""Admin -> Planner: the week's harvest plan against the owner's estimate,
what has been picked, and the cartons the receivers are waiting for.

Three sources, one screen:

- Boord's own crates (HarvestRecord) for what each block has given so far;
- Boord Owner's owner.db, read-only, for the owner's latest estimate of the
  season (kg per tree x trees, per block) - the figure "remaining" is
  measured against;
- Boord Packhouse's packhouse.db, read-only, for the receivers, the packs
  and the delivery notes already loaded.

Only the plan itself (models.HarvestPlan) and the cartons-per-pallet of each
pack (models.PlannerPack) are written, and only to boord.db. The two sibling
databases are opened the way reports.py opens packhouse.db: mode=ro, and a
503 in words rather than a 500 when one is not on this PC - the plan must
still open with the estimate or the receivers missing, saying so.

The arithmetic is in planner.py, where scripts/selftest.py can reach it.
"""
import io
import json
import os
import re
import sqlite3
from datetime import date, datetime, timedelta
from typing import Dict, List, Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field
from sqlalchemy import func
from sqlmodel import Session, select

import planner
from db import DATA_DIR, get_own_supplier_id, get_session
from excel_io import XLSX_MEDIA
from models import Block, HarvestPlan, HarvestRecord, PlannerPack, SystemSetting
from routers.reports import REPORTS_DIR, _packhouse_rows
from security import require_admin_client
from timeutil import day_bounds, to_local

# Every endpoint here is admin-only.
router = APIRouter(prefix="/api/planner", tags=["planner"], dependencies=[Depends(require_admin_client)])

# Boord Owner keeps the owner's estimate in its own owner.db. Beside this repo
# on the Mac and on the farm PC (C:\Boord and C:\BoordOwner); the env var is
# for anywhere else. Opened read-only, like Packhouse's (routers/reports.py).
OWNER_DB_PATH = os.path.abspath(os.environ.get("BOORD_OWNER_DB", os.path.join(
    DATA_DIR, "..", "..", "BoordOwner", "data", "owner.db")))

# A receiver's market group, by the word USA or EU in its market name, else
# local: the same rule Packhouse applies to a pallet's market
# (export_forms.market_group), so the planner's headings match its screens.
EXPORT_MARKET = re.compile(r"\b(usa|eu)\b", re.I)
GROUP_ORDER = {"usa": 0, "eu": 1, "local": 2}
GROUP_LABEL = {"usa": "USA", "eu": "EU", "local": "Local"}


def market_group(name: str) -> str:
    hit = EXPORT_MARKET.search(name or "")
    return hit.group(1).lower() if hit else "local"


# ---------------------------------------------------------------------------
# The figures
# ---------------------------------------------------------------------------

def _season_anchor(session: Session) -> tuple:
    settings = session.exec(select(SystemSetting)).first()
    return ((settings.season_start_month if settings else None) or 1,
            (settings.season_start_day if settings else None) or 1)


def _owner_estimate(season_year: int) -> dict:
    """The owner's latest estimate for the season, per block, from owner.db.
    "Latest" is the version saved or edited last, the one Boord Owner's own
    Estimate tab opens on. Blocks the owner has not put a figure against are
    left out, and the screen shows them without a remaining figure."""
    none = {"available": False, "blocks": {}}
    if not os.path.exists(OWNER_DB_PATH):
        return dict(none, reason="Boord Owner data was not found on this server, so there is no season estimate")
    conn = sqlite3.connect(f"file:{OWNER_DB_PATH}?mode=ro", uri=True, timeout=10)
    conn.row_factory = sqlite3.Row
    try:
        est = conn.execute(
            "SELECT id, name, updated_at FROM yieldestimate WHERE season_year = ? "
            "ORDER BY updated_at DESC, id DESC LIMIT 1", (season_year,)).fetchone()
        if not est:
            return dict(none, reason=f"Boord Owner has no {season_year} estimate yet")
        lines = conn.execute("SELECT block_id, trees, kg_per_tree FROM yieldestimateblock WHERE estimate_id = ?",
                             (est["id"],)).fetchall()
    except sqlite3.OperationalError as e:
        return dict(none, reason=f"Boord Owner data could not be read: {e}")
    finally:
        conn.close()
    blocks = {l["block_id"]: round((l["trees"] or 0) * l["kg_per_tree"], 1)
              for l in lines if l["kg_per_tree"] is not None}
    return {"available": True, "name": est["name"] or "", "updated_at": est["updated_at"],
            "total_kg": round(sum(blocks.values()), 1), "blocks": blocks}


def _block_kg(session: Session, start: date, end: date) -> Dict[str, float]:
    """Net kg (weight less deductions) per block over a span of local days,
    inclusive. Every crate captured, whoever's workers picked it."""
    if end < start:
        return {}
    s, e = day_bounds(start, end)
    rows = session.exec(
        select(HarvestRecord.block_id, func.sum(HarvestRecord.weight_kg - HarvestRecord.deduction_kg))
        .where(HarvestRecord.timestamp >= s, HarvestRecord.timestamp <= e)
        .group_by(HarvestRecord.block_id)).all()
    return {bid: round(kg or 0.0, 1) for bid, kg in rows if bid}


def _week_actuals(session: Session, week_start: date) -> Dict[str, List[float]]:
    """Net kg per block for each day of the week, Monday first, so the plan
    shows what was really picked under each planned figure as the week runs."""
    s, e = day_bounds(week_start, week_start + timedelta(days=6))
    out: Dict[str, List[float]] = {}
    for r in session.exec(select(HarvestRecord).where(HarvestRecord.timestamp >= s,
                                                     HarvestRecord.timestamp <= e)).all():
        if not r.block_id:
            continue
        local_day = to_local(r.timestamp).date()
        idx = (local_day - week_start).days
        if 0 <= idx <= 6:
            days = out.setdefault(r.block_id, [0.0] * 7)
            days[idx] += r.weight_kg - r.deduction_kg
    return {bid: [round(x, 1) for x in days] for bid, days in out.items()}


def _packhouse(season_start: date, week_start: date) -> dict:
    """Receivers, packs and what went out on delivery notes before this week:
    cartons per receiver and pack for the season to date and for last week,
    plus the fruit kg delivered so far (for the loss share). Receivers and
    packs retired in Packhouse still appear when a note this season used
    them, so their YTD stays visible."""
    try:
        receivers = _packhouse_rows("SELECT id, name, market, active, sort FROM receiver ORDER BY sort, name")
        packs = _packhouse_rows("SELECT name, carton_kg, active, sort FROM packaging ORDER BY sort, name")
        notes = _packhouse_rows(
            "SELECT receiver_id, loaded_on, pallets FROM deliverynote WHERE removed = 0 "
            "AND loaded_on >= ? AND loaded_on < ?", (season_start.isoformat(), week_start.isoformat()))
    except HTTPException as e:
        return {"available": False, "reason": e.detail, "receivers": [], "packs": [],
                "delivered": {}, "delivered_kg": 0.0}
    last_week_start = (week_start - timedelta(days=7)).isoformat()
    delivered: Dict[str, Dict[str, dict]] = {}
    used_receivers, used_packs = set(), {}
    delivered_kg = 0.0
    for n in notes:
        rid = str(n["receiver_id"] or 0)
        used_receivers.add(rid)
        in_last_week = (n["loaded_on"] or "") >= last_week_start
        for p in json.loads(n["pallets"] or "[]"):
            for ln in p.get("lines", []):
                name = ln.get("packaging_name") or ""
                cartons = int(ln.get("cartons") or 0)
                delivered_kg += cartons * (ln.get("carton_kg") or 0)
                used_packs.setdefault(name, ln.get("carton_kg") or 0)
                cell = delivered.setdefault(rid, {}).setdefault(name, {"ytd": 0, "last_week": 0})
                cell["ytd"] += cartons
                if in_last_week:
                    cell["last_week"] += cartons
    out_receivers = [{"id": r["id"], "name": r["name"], "market": r["market"] or "",
                      "group": market_group(r["market"]), "active": bool(r["active"])}
                     for r in receivers if r["active"] or str(r["id"]) in used_receivers]
    if "0" in used_receivers:
        out_receivers.append({"id": 0, "name": "No receiver", "market": "", "group": "local", "active": False})
    out_receivers.sort(key=lambda r: (GROUP_ORDER[r["group"]], r["name"].lower()))
    out_packs = [{"name": p["name"], "carton_kg": p["carton_kg"], "active": bool(p["active"])}
                 for p in packs if p["active"] or p["name"] in used_packs]
    known = {p["name"] for p in out_packs}
    for name, kg in used_packs.items():
        if name and name not in known:
            out_packs.append({"name": name, "carton_kg": kg, "active": False})
    out_packs.sort(key=lambda p: planner.pack_sort_key(p["name"]))
    return {"available": True, "receivers": out_receivers, "packs": out_packs, "delivered": delivered,
            "delivered_kg": round(delivered_kg, 1)}


def _seasons(session: Session, month: int, day: int, current: int) -> List[int]:
    """Every season with a crate or a plan in it, plus the current one."""
    years = {current}
    first = session.exec(select(func.min(HarvestRecord.timestamp))).one()
    if first:
        years.add(planner.season_year_for(to_local(first).date(), month, day))
    for y in session.exec(select(HarvestPlan.season_year).distinct()).all():
        years.add(y)
    lo, hi = min(years), max(years)
    return list(range(lo, hi + 1))


def build_week(session: Session, week_start: date) -> dict:
    """Everything the Planner screen, the spreadsheet and the PDF need for
    one week, figures and saved plan together, with the derived totals."""
    month, day = _season_anchor(session)
    season_year = planner.season_year_for(week_start, month, day)
    season_start, season_end = planner.season_bounds(season_year, month, day)
    today = datetime.now().date()
    week_end = week_start + timedelta(days=6)

    own_id = get_own_supplier_id(session)
    blocks = [b for b in session.exec(select(Block).where(Block.active == True)).all()  # noqa: E712
              if b.supplier_id is None or b.supplier_id == own_id]
    blocks.sort(key=lambda b: planner.block_sort_key(b.id))

    estimate = _owner_estimate(season_year)
    ytd = _block_kg(session, season_start, week_start - timedelta(days=1))
    last_week = _block_kg(session, week_start - timedelta(days=7), week_start - timedelta(days=1))
    actuals = _week_actuals(session, week_start)
    ph = _packhouse(season_start, week_start)

    plan = session.get(HarvestPlan, week_start)
    saved_blocks = {p["block_id"]: p for p in json.loads(plan.blocks or "[]")} if plan else {}

    out_blocks = []
    for b in blocks:
        p = saved_blocks.get(b.id, {})
        out_blocks.append({
            "block_id": b.id, "name": b.name or b.id, "variety": b.variety,
            "estimate_kg": estimate["blocks"].get(b.id),
            "ytd_kg": ytd.get(b.id, 0.0), "last_week_kg": last_week.get(b.id, 0.0),
            "actual": actuals.get(b.id, [0.0] * 7),
            "on": bool(p.get("on", False)),  # a week not yet planned starts with every block off
            "available_kg": p.get("available_kg"),
            "days": [d for d in (p.get("days") or []) if isinstance(d, int) and 0 <= d <= 6],
            "fixed": {str(k): v for k, v in (p.get("fixed") or {}).items()},
        })

    settings = {pp.pack: pp for pp in session.exec(select(PlannerPack)).all()}
    packs = [dict(pk, cartons_per_pallet=settings[pk["name"]].cartons_per_pallet if pk["name"] in settings else 0,
                  units_per_carton=settings[pk["name"]].units_per_carton if pk["name"] in settings
                  else default_units_per_carton(pk["name"]))
             for pk in ph["packs"]]

    # The season's own loss share: everything picked before this week against
    # the fruit kg loaded before it. Juice bins are not fruit, so they count
    # as loss, which is what the office means by the word.
    harvested_total = round(sum(ytd.values()), 1)
    weeks = planner.season_weeks(season_year, month, day)
    plan_weeks = set(session.exec(
        select(HarvestPlan.week_start).where(HarvestPlan.season_year == season_year)).all())
    view = {
        "week_start": week_start.isoformat(), "week_end": week_end.isoformat(),
        "iso_week": week_start.isocalendar()[1], "iso_year": week_start.isocalendar()[0],
        "season_year": season_year, "season_start": season_start.isoformat(), "season_end": season_end.isoformat(),
        "seasons": _seasons(session, month, day, planner.season_year_for(today, month, day)),
        "weeks": [{"week_start": w.isoformat(), "iso_week": w.isocalendar()[1], "has_plan": w in plan_weeks,
                   "current": w <= today <= w + timedelta(days=6)} for w in weeks],
        "today": today.isoformat(),
        "estimate": {k: v for k, v in estimate.items() if k != "blocks"},
        "packhouse": {"available": ph["available"], "reason": ph.get("reason")},
        "blocks": out_blocks,
        "receivers": ph["receivers"],
        "packs": packs,
        "delivered": ph["delivered"],
        "plan_deliveries": json.loads(plan.deliveries or "{}") if plan else {},
        "loss": {
            "harvested_kg": harvested_total, "delivered_kg": ph["delivered_kg"],
            "season_pct": planner.loss_pct(harvested_total, ph["delivered_kg"]) if ph["available"] else None,
            "override_pct": plan.loss_pct if plan else None,
        },
        "updated_at": plan.updated_at.isoformat() + "Z" if plan else None,
    }
    view["derived"] = planner.derive(view)
    return view


def default_units_per_carton(pack_name: str) -> int:
    """Until the office types it: Packhouse counts Pre-pac in punnets, four
    to a carton; every other pack's unit is the carton itself."""
    name = (pack_name or "").lower().replace("-", "").replace(" ", "")
    return 4 if "prepac" in name else 1


def _monday(d: date) -> date:
    return planner.monday_of(d)


# ---------------------------------------------------------------------------
# Endpoints
# ---------------------------------------------------------------------------

@router.get("/week")
def get_week(week_start: Optional[date] = Query(default=None), season: Optional[int] = Query(default=None),
             session: Session = Depends(get_session)):
    """The plan for a week: by its Monday, or this week's when none is given.
    With `season` alone, the season's week that holds today, or its first."""
    if week_start is None:
        today = datetime.now().date()
        if season is not None:
            month, day = _season_anchor(session)
            weeks = planner.season_weeks(season, month, day)
            week_start = next((w for w in weeks if w <= today <= w + timedelta(days=6)), weeks[0])
        else:
            week_start = today
    return build_week(session, _monday(week_start))


class BlockPlanIn(BaseModel):
    block_id: str
    on: bool = True
    available_kg: Optional[float] = Field(default=None, ge=0)
    days: List[int] = []
    fixed: Dict[str, float] = {}


class PlanIn(BaseModel):
    loss_pct: Optional[float] = Field(default=None, ge=-100, le=100)
    blocks: List[BlockPlanIn] = []
    deliveries: Dict[str, Dict[str, int]] = {}


@router.put("/week/{week_start}")
def save_week(week_start: date, body: PlanIn, session: Session = Depends(get_session)):
    """Save the week's plan whole. One plan per week, the last save wins:
    the screen saves as the office types, so there is nothing to merge."""
    week_start = _monday(week_start)
    month, day = _season_anchor(session)
    blocks = []
    for b in body.blocks:
        days = sorted({d for d in b.days if 0 <= d <= 6})
        fixed = {k: round(float(v), 1) for k, v in b.fixed.items()
                 if k.isdigit() and int(k) in days and v is not None and v >= 0}
        blocks.append({"block_id": b.block_id, "on": b.on, "available_kg": b.available_kg,
                       "days": days, "fixed": fixed})
    deliveries = {rid: {pack: int(c) for pack, c in cells.items() if c}
                  for rid, cells in body.deliveries.items()}
    deliveries = {rid: cells for rid, cells in deliveries.items() if cells}
    session.merge(HarvestPlan(
        week_start=week_start, season_year=planner.season_year_for(week_start, month, day),
        blocks=json.dumps(blocks), deliveries=json.dumps(deliveries), loss_pct=body.loss_pct,
        updated_at=datetime.utcnow(),
    ))
    session.commit()
    return build_week(session, week_start)


class PackIn(BaseModel):
    pack: str
    cartons_per_pallet: int = Field(default=0, ge=0)
    units_per_carton: int = Field(default=1, ge=1)


@router.put("/packs")
def save_packs(body: List[PackIn], session: Session = Depends(get_session)):
    """Cartons per pallet and Packhouse units per carton for each pack. Zero
    cartons per pallet means not known: pallets then show as a dash for that
    pack rather than a division by nothing. Pallets are counted whole,
    rounded up (planner.derive)."""
    for p in body:
        if p.pack.strip():
            session.merge(PlannerPack(pack=p.pack.strip(), cartons_per_pallet=p.cartons_per_pallet,
                                      units_per_carton=p.units_per_carton))
    session.commit()
    return {"ok": True}


# ---------------------------------------------------------------------------
# Exports: the whole page as one sheet, or one landscape PDF
# ---------------------------------------------------------------------------

def _title(view: dict) -> str:
    ws, we = date.fromisoformat(view["week_start"]), date.fromisoformat(view["week_end"])
    return (f"Harvest plan - {view['season_year']} season, week {view['iso_week']} "
            f"({ws.strftime('%d %b')} to {we.strftime('%d %b %Y')})")


def _num(v):
    """A whole number for a cell, or blank: the exports say nothing where the
    screen shows a dash, rather than a 0 that reads as a measurement."""
    if v is None:
        return ""
    return int(round(v))


def _pct_text(v) -> str:
    return "" if v is None else f"{int(round(v))}%"


def _harvest_rows(view: dict, compact: bool = False) -> tuple:
    """Headers and rows of the harvest section, blocks that are on only, with
    a totals row. The sheet gets everything the screen shows; the PDF
    (`compact`) is the sheet pinned up for the week - block, the days and
    the total, and only blocks with something planned - because the office
    reads estimate, YTD and remaining on the screen, not on paper."""
    d = view["derived"]
    derived = {b["block_id"]: b for b in d["blocks"]}
    detail = ["Estimate kg", "YTD kg", "YTD %", "Last wk kg", "Last wk %", "Remaining kg", "Remaining %",
              "Available kg"]
    headers = ["Block"] + ([] if compact else detail) + planner.WEEKDAYS + ["Total kg"]
    rows = []
    for b in view["blocks"]:
        x = derived[b["block_id"]]
        if not b["on"] or (compact and not x["total_kg"]):
            continue
        row = [b["name"]]
        if not compact:
            row += [_num(b["estimate_kg"]), _num(b["ytd_kg"]), _pct_text(x["ytd_pct"]),
                    _num(b["last_week_kg"]), _pct_text(x["last_week_pct"]), _num(x["remaining_kg"]),
                    _pct_text(x["remaining_pct"]), _num(x["available_kg"])]
        rows.append(row + [_num(x["per_day"].get(str(i))) for i in range(7)] + [_num(x["total_kg"])])
    on = [b for b in view["blocks"] if b["on"]]
    total = ["Total"]
    if not compact:
        total += [
            _num(sum(b["estimate_kg"] or 0 for b in on)) if any(b["estimate_kg"] is not None for b in on) else "",
            _num(sum(b["ytd_kg"] for b in on)), "", _num(sum(b["last_week_kg"] for b in on)), "",
            _num(sum((derived[b["block_id"]]["remaining_kg"] or 0) for b in on)), "",
            _num(sum(derived[b["block_id"]]["available_kg"] for b in on))]
    total += [_num(t) for t in d["day_totals"]] + [_num(d["harvest_kg"])]
    return headers, rows, total


def _delivery_rows(view: dict, compact: bool = False) -> tuple:
    """The delivery section. The sheet has YTD and last week beside this
    week's cartons for each pack; the PDF (`compact`) has this week's only.
    The receiver's Cartons, Kg and Pallets come right after its name, before
    the packs, as on the screen: with many packs the row runs wide and the
    totals would otherwise sit far from the receiver they belong to."""
    d = view["derived"]
    packs = view["packs"]
    headers = ["Market", "Receiver", "Cartons", "Kg", "Pallets"]
    for p in packs:
        headers += ([p["name"]] if compact
                    else [f"{p['name']} YTD", f"{p['name']} last wk", f"{p['name']} this wk"])
    rows = []
    derived = {str(r["id"]): r for r in d["receivers"]}
    for r in view["receivers"]:
        x = derived[str(r["id"])]
        row = [GROUP_LABEL[r["group"]], r["name"], x["cartons"], _num(x["kg"]), x["pallets"]]
        for p in packs:
            if not compact:
                done = (view["delivered"].get(str(r["id"])) or {}).get(p["name"]) or {}
                upc = p.get("units_per_carton") or 1
                row += [planner.units_to_cartons(done.get("ytd", 0), upc),
                        planner.units_to_cartons(done.get("last_week", 0), upc)]
            row.append(x["cells"][p["name"]]["cartons"])
        rows.append(row)
    total = ["Total", "", d["delivery_cartons"], _num(d["delivery_kg"]), d["delivery_pallets"]]
    for p in packs:
        if not compact:
            cells = [(view["delivered"].get(str(r["id"])) or {}).get(p["name"]) or {} for r in view["receivers"]]
            upc = p.get("units_per_carton") or 1
            total += [planner.units_to_cartons(sum(c.get("ytd", 0) for c in cells), upc),
                      planner.units_to_cartons(sum(c.get("last_week", 0) for c in cells), upc)]
        total.append(d["pack_totals"][p["name"]]["cartons"])
    return headers, rows, total


def _grand_rows(view: dict, compact: bool = False) -> list:
    """The grand totals. The sheet says where the loss % came from; the PDF
    (`compact`) just gives the figure."""
    d = view["derived"]
    loss = view["loss"]
    source = ("typed" if loss.get("override_pct") is not None
              else "season to date" if loss.get("season_pct") is not None else "none yet")
    return [
        ["Planned harvest kg", _num(d["harvest_kg"])],
        ["Juice and loss %" if compact else f"Juice and loss % ({source})", _num(d["loss_pct"])],
        ["Deliverable kg", _num(d["deliverable_kg"])],
        ["Planned delivery kg", _num(d["delivery_kg"])],
        ["Surplus kg" if d["balance_kg"] >= 0 else "Short kg", _num(abs(d["balance_kg"]))],
    ]


def _keep_copy(filename: str, data: bytes) -> None:
    # Every export is also kept in data/reports/, like the Reports tab's, so
    # the backup that covers data/ has it too.
    with open(os.path.join(REPORTS_DIR, filename), "wb") as f:
        f.write(data)


@router.get("/week/{week_start}/export.xlsx")
def export_xlsx(week_start: date, session: Session = Depends(get_session)):
    view = build_week(session, _monday(week_start))
    wb = Workbook()
    ws = wb.active
    ws.title = f"Week {view['iso_week']}"
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="E2E8F0")

    def section(title, headers, rows, total):
        ws.append([title])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12)
        ws.append(headers)
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=ws.max_row, column=c)
            cell.font, cell.fill = bold, head_fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for r in rows:
            ws.append(r)
        if total:
            ws.append(total)
            for c in range(1, len(total) + 1):
                ws.cell(row=ws.max_row, column=c).font = bold
        ws.append([])

    ws.append([_title(view)])
    ws.cell(row=1, column=1).font = Font(bold=True, size=14)
    if view["estimate"].get("available"):
        ws.append([f"Estimate: {view['estimate'].get('name') or 'latest'} (Boord Owner)"])
    else:
        ws.append([view["estimate"].get("reason", "")])
    if not view["packhouse"]["available"]:
        ws.append([view["packhouse"].get("reason", "")])
    ws.append([])
    section("Harvest this week", *_harvest_rows(view))
    section("Delivery this week (cartons; Packhouse's units per carton as set on the Planner)", *_delivery_rows(view))
    section("Grand totals", ["", ""], _grand_rows(view), None)
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None and cell.row > 3:
                widths[cell.column] = max(widths.get(cell.column, 0), min(len(str(cell.value)), 28))
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(8, w + 2)
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    filename = f"Harvest_Plan_{view['season_year']}_week_{view['iso_week']:02d}.xlsx"
    _keep_copy(filename, data)
    return Response(content=data, media_type=XLSX_MEDIA,
                    headers={"Content-Disposition": f'attachment; filename="{filename}"'})


def _latin(s) -> str:
    """fpdf2's built-in Helvetica knows latin-1 only. Block and receiver names
    are Afrikaans and fit; anything else becomes a question mark rather than
    an error that stops the whole export."""
    return str(s).encode("latin-1", "replace").decode("latin-1")


@router.get("/week/{week_start}/export.pdf")
def export_pdf(week_start: date, session: Session = Depends(get_session)):
    view = build_week(session, _monday(week_start))
    pdf = FPDF(orientation="L", format="A4")
    pdf.set_auto_page_break(auto=True, margin=10)
    pdf.set_margins(10, 10, 10)
    pdf.add_page()
    width = pdf.w - 20

    def text(s, size=9, style="", h=5):
        pdf.set_font("Helvetica", style, size)
        pdf.cell(0, h, _latin(s), new_x="LMARGIN", new_y="NEXT")

    def table(headers, rows, total, first_col=38):
        n = len(headers)
        rest = (width - first_col) / max(1, n - 1)
        widths = [first_col] + [rest] * (n - 1)
        # Headings wrap inside their cell (two lines at most for the widths
        # here) rather than running into the next one: "Remaining %" and
        # "4.5 kg XL last wk" are wider than a day column.
        pdf.set_font("Helvetica", "B", 6.5)
        pdf.set_fill_color(226, 232, 240)
        y0, x = pdf.get_y(), pdf.l_margin
        head_h = 9
        for h, w in zip(headers, widths):
            pdf.set_xy(x, y0)
            pdf.cell(w, head_h, "", border=1, fill=True)
            pdf.set_xy(x, y0 + 0.5)
            pdf.multi_cell(w, 4, _latin(h), border=0, align="C", max_line_height=4)
            x += w
        pdf.set_xy(pdf.l_margin, y0 + head_h)
        pdf.set_font("Helvetica", "", 8)
        for r in rows:
            if pdf.get_y() > pdf.h - 20:
                pdf.add_page()
            for i, (v, w) in enumerate(zip(r, widths)):
                pdf.cell(w, 6, _latin(v), border=1, align="L" if i == 0 else "R")
            pdf.ln(6)
        if total:
            pdf.set_font("Helvetica", "B", 8)
            for i, (v, w) in enumerate(zip(total, widths)):
                pdf.cell(w, 6, _latin(v), border=1, align="L" if i == 0 else "R")
            pdf.ln(6)
        pdf.ln(3)

    text(_title(view), 13, "B", 8)
    if view["estimate"].get("available"):
        text(f"Estimate: {view['estimate'].get('name') or 'latest'} (Boord Owner)", 8)
    else:
        text(view["estimate"].get("reason", ""), 8)
    if not view["packhouse"]["available"]:
        text(view["packhouse"].get("reason", ""), 8)
    pdf.ln(2)
    text("Harvest this week (kg)", 11, "B", 7)
    table(*_harvest_rows(view, compact=True), first_col=50)
    text("Delivery this week (cartons)", 11, "B", 7)
    h, r, t = _delivery_rows(view, compact=True)
    table(h, r, t, first_col=28)
    text("Grand totals", 11, "B", 7)
    pdf.set_font("Helvetica", "", 10)
    for label, value in _grand_rows(view, compact=True):
        short = label.startswith("Short")
        if label.startswith(("Short", "Surplus")):
            pdf.set_text_color(185, 28, 28) if short else pdf.set_text_color(21, 128, 61)
            pdf.set_font("Helvetica", "B", 10)
        pdf.cell(70, 7, _latin(label), border=1)
        pdf.cell(40, 7, _latin(value), border=1, align="R", new_x="LMARGIN", new_y="NEXT")
        pdf.set_text_color(0, 0, 0)
        pdf.set_font("Helvetica", "", 10)
    data = bytes(pdf.output())
    filename = f"Harvest_Plan_{view['season_year']}_week_{view['iso_week']:02d}.pdf"
    _keep_copy(filename, data)
    return Response(content=data, media_type="application/pdf",
                    headers={"Content-Disposition": f'attachment; filename="{filename}"'})
